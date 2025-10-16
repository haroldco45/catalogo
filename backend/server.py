from fastapi import FastAPI, APIRouter, HTTPException, Query
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timedelta
import pytz
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.responses import StreamingResponse
import requests


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# Colombia timezone
COLOMBIA_TZ = pytz.timezone('America/Bogota')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# WhatsApp Configuration
WHATSAPP_PHONE = "573217366758"  # Formato internacional: 57 + número
WHATSAPP_API_PROVIDER = os.environ.get('WHATSAPP_API_PROVIDER', 'callmebot')
WHAPI_TOKEN = os.environ.get('WHAPI_TOKEN', '')
WHATSAPP_BUSINESS_APIKEY = os.environ.get('WHATSAPP_BUSINESS_APIKEY', '')
WHATSAPP_CLIENTS_APIKEYS = os.environ.get('WHATSAPP_CLIENTS_APIKEYS', '')

# ==================== MODELS ====================

# Product Categories
CATEGORIES = [
    "Galletas",
    "Chocolates", 
    "Helados",
    "Cafés",
    "Cárnicos",
    "Cuidado Personal",
    "Limpieza del Hogar",
    "Alimentos Básicos",
    "Delicatessen",
    "Moda"
]

class Product(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    description: Optional[str] = ""
    category: str
    price: float
    stock: int
    min_stock: int = 10
    supplier: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(COLOMBIA_TZ))
    updated_at: datetime = Field(default_factory=lambda: datetime.now(COLOMBIA_TZ))

class ProductCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    category: str
    price: float
    stock: int
    min_stock: int = 10
    supplier: Optional[str] = ""

class ProductUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    category: Optional[str] = None
    price: Optional[float] = None
    stock: Optional[int] = None
    min_stock: Optional[int] = None
    supplier: Optional[str] = None

class Customer(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    phone: str
    address: Optional[str] = ""
    email: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(COLOMBIA_TZ))

class CustomerCreate(BaseModel):
    name: str
    phone: str
    address: Optional[str] = ""
    email: Optional[str] = ""

class CustomerUpdate(BaseModel):
    name: Optional[str] = None
    phone: Optional[str] = None
    address: Optional[str] = None
    email: Optional[str] = None

class SaleItem(BaseModel):
    product_id: str
    product_name: str
    quantity: int
    price: float
    subtotal: float

class Sale(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    customer_id: Optional[str] = None
    customer_name: str
    items: List[SaleItem]
    total: float
    payment_method: str = "Efectivo"
    status: str = "Completado"
    notes: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(COLOMBIA_TZ))

class SaleCreate(BaseModel):
    customer_id: Optional[str] = None
    customer_name: str
    items: List[SaleItem]
    total: float
    payment_method: str = "Efectivo"
    notes: Optional[str] = ""

# ==================== HELPER FUNCTIONS ====================

def datetime_to_str(dt):
    """Convert datetime to ISO string for MongoDB"""
    if isinstance(dt, datetime):
        return dt.isoformat()
    return dt

def str_to_datetime(dt_str):
    """Convert ISO string to datetime"""
    if isinstance(dt_str, str):
        return datetime.fromisoformat(dt_str)
    return dt_str

def normalize_phone_number(phone: str) -> str:
    """Normalize phone number to Colombia international format (57XXXXXXXXXX)"""
    if not phone:
        return phone
    
    # Remove all non-digit characters
    cleaned = ''.join(filter(str.isdigit, phone))
    
    # If starts with 57, return as is
    if cleaned.startswith('57'):
        return cleaned
    
    # If starts with 0, remove it (Colombian mobile format)
    if cleaned.startswith('0'):
        cleaned = cleaned[1:]
    
    # Add Colombia code (57)
    return f"57{cleaned}"

async def send_whatsapp_message(phone: str, message: str):
    """Send WhatsApp message using Whapi.cloud or CallMeBot API"""
    try:
        # Use Whapi.cloud if configured
        if WHATSAPP_API_PROVIDER == 'whapi' and WHAPI_TOKEN:
            url = "https://gate.whapi.cloud/messages/text"
            headers = {
                "Authorization": f"Bearer {WHAPI_TOKEN}",
                "Content-Type": "application/json"
            }
            data = {
                "to": phone,
                "body": message
            }
            
            response = requests.post(url, json=data, headers=headers, timeout=10)
            
            if response.status_code in [200, 201]:
                logger.info(f"✅ WhatsApp ENVIADO a {phone} (Whapi.cloud)")
                return True
            else:
                logger.error(f"❌ Error Whapi.cloud a {phone}: {response.status_code} - {response.text}")
                return False
        
        # Fallback to CallMeBot
        else:
            # Get API key for this phone number
            api_key = None
            
            # Check if it's the business phone
            if phone == WHATSAPP_PHONE and WHATSAPP_BUSINESS_APIKEY:
                api_key = WHATSAPP_BUSINESS_APIKEY
            else:
                # Check client API keys
                if WHATSAPP_CLIENTS_APIKEYS:
                    clients_keys = {}
                    for pair in WHATSAPP_CLIENTS_APIKEYS.split(','):
                        if ':' in pair:
                            client_phone, client_key = pair.strip().split(':')
                            clients_keys[client_phone] = client_key
                    api_key = clients_keys.get(phone)
            
            # If no API key, just log (for testing without real API)
            if not api_key:
                logger.info(f"WhatsApp to {phone} (SIN API KEY - SOLO LOG): {message}")
                return True
            
            # Send via CallMeBot API
            import urllib.parse
            encoded_message = urllib.parse.quote(message)
            url = f"https://api.callmebot.com/whatsapp.php?phone={phone}&text={encoded_message}&apikey={api_key}"
            
            response = requests.get(url, timeout=10)
            
            if response.status_code == 200:
                logger.info(f"✅ WhatsApp ENVIADO a {phone} (CallMeBot)")
                return True
            else:
                logger.error(f"❌ Error enviando WhatsApp a {phone}: {response.status_code} - {response.text}")
                return False
            
    except Exception as e:
        logger.error(f"❌ Error enviando WhatsApp a {phone}: {e}")
        return False

# ==================== PRODUCTS API ====================

@api_router.get("/categories")
async def get_categories():
    """Get all available categories"""
    return {"categories": CATEGORIES}

@api_router.post("/products", response_model=Product)
async def create_product(product_data: ProductCreate):
    """Create a new product"""
    product = Product(**product_data.model_dump())
    doc = product.model_dump()
    doc['created_at'] = datetime_to_str(doc['created_at'])
    doc['updated_at'] = datetime_to_str(doc['updated_at'])
    
    await db.products.insert_one(doc)
    return product

@api_router.get("/products", response_model=List[Product])
async def get_products(
    category: Optional[str] = None,
    search: Optional[str] = None,
    low_stock: bool = False
):
    """Get all products with optional filters"""
    query = {}
    
    if category:
        query['category'] = category
    
    if search:
        query['$or'] = [
            {'name': {'$regex': search, '$options': 'i'}},
            {'description': {'$regex': search, '$options': 'i'}}
        ]
    
    if low_stock:
        query['$expr'] = {'$lte': ['$stock', '$min_stock']}
    
    products = await db.products.find(query, {"_id": 0}).to_list(1000)
    
    for product in products:
        product['created_at'] = str_to_datetime(product['created_at'])
        product['updated_at'] = str_to_datetime(product['updated_at'])
    
    return products

@api_router.get("/products/{product_id}", response_model=Product)
async def get_product(product_id: str):
    """Get a specific product"""
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    product['created_at'] = str_to_datetime(product['created_at'])
    product['updated_at'] = str_to_datetime(product['updated_at'])
    
    return product

@api_router.put("/products/{product_id}", response_model=Product)
async def update_product(product_id: str, product_data: ProductUpdate):
    """Update a product"""
    update_data = {k: v for k, v in product_data.model_dump().items() if v is not None}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="No hay datos para actualizar")
    
    update_data['updated_at'] = datetime_to_str(datetime.now(COLOMBIA_TZ))
    
    result = await db.products.update_one(
        {"id": product_id},
        {"$set": update_data}
    )
    
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    updated_product = await db.products.find_one({"id": product_id}, {"_id": 0})
    updated_product['created_at'] = str_to_datetime(updated_product['created_at'])
    updated_product['updated_at'] = str_to_datetime(updated_product['updated_at'])
    
    return updated_product

@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str):
    """Delete a product"""
    result = await db.products.delete_one({"id": product_id})
    
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    return {"message": "Producto eliminado exitosamente"}

# ==================== CUSTOMERS API ====================

@api_router.post("/customers", response_model=Customer)
async def create_customer(customer_data: CustomerCreate):
    """Create a new customer"""
    customer_dict = customer_data.model_dump()
    
    # Normalize phone number to Colombia format (+57)
    if customer_dict.get('phone'):
        customer_dict['phone'] = normalize_phone_number(customer_dict['phone'])
    
    customer = Customer(**customer_dict)
    doc = customer.model_dump()
    doc['created_at'] = datetime_to_str(doc['created_at'])
    
    await db.customers.insert_one(doc)
    return customer

@api_router.get("/customers", response_model=List[Customer])
async def get_customers(search: Optional[str] = None):
    """Get all customers"""
    query = {}
    
    if search:
        query['$or'] = [
            {'name': {'$regex': search, '$options': 'i'}},
            {'phone': {'$regex': search, '$options': 'i'}}
        ]
    
    customers = await db.customers.find(query, {"_id": 0}).to_list(1000)
    
    for customer in customers:
        customer['created_at'] = str_to_datetime(customer['created_at'])
    
    return customers

@api_router.get("/customers/{customer_id}", response_model=Customer)
async def get_customer(customer_id: str):
    """Get a specific customer"""
    customer = await db.customers.find_one({"id": customer_id}, {"_id": 0})
    
    if not customer:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    
    customer['created_at'] = str_to_datetime(customer['created_at'])
    
    return customer

@api_router.put("/customers/{customer_id}", response_model=Customer)
async def update_customer(customer_id: str, customer_data: CustomerUpdate):
    """Update a customer"""
    update_data = {k: v for k, v in customer_data.model_dump().items() if v is not None}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="No hay datos para actualizar")
    
    # Normalize phone number if provided
    if 'phone' in update_data and update_data['phone']:
        update_data['phone'] = normalize_phone_number(update_data['phone'])
    
    result = await db.customers.update_one(
        {"id": customer_id},
        {"$set": update_data}
    )
    
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    
    updated_customer = await db.customers.find_one({"id": customer_id}, {"_id": 0})
    updated_customer['created_at'] = str_to_datetime(updated_customer['created_at'])
    
    return updated_customer

@api_router.delete("/customers/{customer_id}")
async def delete_customer(customer_id: str):
    """Delete a customer"""
    result = await db.customers.delete_one({"id": customer_id})
    
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    
    return {"message": "Cliente eliminado exitosamente"}

# ==================== SALES API ====================

@api_router.post("/sales", response_model=Sale)
async def create_sale(sale_data: SaleCreate):
    """Create a new sale and update inventory"""
    sale = Sale(**sale_data.model_dump())
    
    # Update product stock
    for item in sale.items:
        product = await db.products.find_one({"id": item.product_id})
        if product:
            new_stock = product['stock'] - item.quantity
            await db.products.update_one(
                {"id": item.product_id},
                {"$set": {"stock": new_stock}}
            )
            
            # Check if stock is low and send WhatsApp alert
            if new_stock <= product.get('min_stock', 10):
                message = f"⚠️ ALERTA NOVAVENTA: El producto '{item.product_name}' tiene stock bajo ({new_stock} unidades)"
                await send_whatsapp_message(WHATSAPP_PHONE, message)
    
    # Save sale
    doc = sale.model_dump()
    doc['created_at'] = datetime_to_str(doc['created_at'])
    await db.sales.insert_one(doc)
    
    # Send sale notification to business owner (3217366758)
    items_text = "\n".join([f"- {item.product_name} x{item.quantity}: ${item.subtotal:,.0f}" for item in sale.items])
    business_message = f"✅ NUEVA VENTA NOVAVENTA\nCliente: {sale.customer_name}\nTotal: ${sale.total:,.0f} COP\n\nProductos:\n{items_text}"
    await send_whatsapp_message(WHATSAPP_PHONE, business_message)
    
    # Send receipt/confirmation to customer via WhatsApp
    customer_phone = None
    if sale.customer_id:
        # Get customer phone from database
        customer = await db.customers.find_one({"id": sale.customer_id})
        if customer and customer.get('phone'):
            customer_phone = customer['phone']
    
    if customer_phone:
        # Format receipt for customer
        sale_date = datetime.now(COLOMBIA_TZ).strftime("%d/%m/%Y %H:%M")
        receipt_items = "\n".join([
            f"  {item.product_name}\n  {item.quantity} x ${item.price:,.0f} = ${item.subtotal:,.0f}"
            for item in sale.items
        ])
        
        customer_message = f"""
🧾 COMPROBANTE DE COMPRA
━━━━━━━━━━━━━━━━━━━━━━
NOVAVENTA
Fecha: {sale_date}

👤 Cliente: {sale.customer_name}

📦 PRODUCTOS:
{receipt_items}

━━━━━━━━━━━━━━━━━━━━━━
💰 TOTAL: ${sale.total:,.0f} COP
💳 Método de Pago: {sale.payment_method}

¡Gracias por tu compra! 🎉
Para consultas: {WHATSAPP_PHONE}
"""
        await send_whatsapp_message(customer_phone, customer_message.strip())
        logger.info(f"Comprobante enviado al cliente {sale.customer_name} ({customer_phone})")
    else:
        logger.info(f"Cliente sin teléfono registrado, no se envió comprobante")
    
    return sale

@api_router.get("/sales", response_model=List[Sale])
async def get_sales(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    customer_id: Optional[str] = None
):
    """Get all sales with optional filters"""
    query = {}
    
    if customer_id:
        query['customer_id'] = customer_id
    
    if start_date or end_date:
        date_query = {}
        if start_date:
            date_query['$gte'] = start_date
        if end_date:
            date_query['$lte'] = end_date
        query['created_at'] = date_query
    
    sales = await db.sales.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    
    for sale in sales:
        sale['created_at'] = str_to_datetime(sale['created_at'])
    
    return sales

@api_router.get("/sales/{sale_id}", response_model=Sale)
async def get_sale(sale_id: str):
    """Get a specific sale"""
    sale = await db.sales.find_one({"id": sale_id}, {"_id": 0})
    
    if not sale:
        raise HTTPException(status_code=404, detail="Venta no encontrada")
    
    sale['created_at'] = str_to_datetime(sale['created_at'])
    
    return sale

# ==================== DASHBOARD & STATS API ====================

@api_router.get("/dashboard/stats")
async def get_dashboard_stats():
    """Get dashboard statistics"""
    now = datetime.now(COLOMBIA_TZ)
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_start = today_start - timedelta(days=today_start.weekday())
    month_start = today_start.replace(day=1)
    
    # Today's sales
    today_sales = await db.sales.find({
        "created_at": {"$gte": datetime_to_str(today_start)}
    }).to_list(1000)
    today_total = sum(sale['total'] for sale in today_sales)
    
    # Week's sales
    week_sales = await db.sales.find({
        "created_at": {"$gte": datetime_to_str(week_start)}
    }).to_list(1000)
    week_total = sum(sale['total'] for sale in week_sales)
    
    # Month's sales
    month_sales = await db.sales.find({
        "created_at": {"$gte": datetime_to_str(month_start)}
    }).to_list(1000)
    month_total = sum(sale['total'] for sale in month_sales)
    
    # Total products
    total_products = await db.products.count_documents({})
    
    # Low stock products
    low_stock_products = await db.products.find({
        "$expr": {"$lte": ["$stock", "$min_stock"]}
    }).to_list(100)
    
    # Total customers
    total_customers = await db.customers.count_documents({})
    
    return {
        "today": {
            "sales": len(today_sales),
            "total": today_total
        },
        "week": {
            "sales": len(week_sales),
            "total": week_total
        },
        "month": {
            "sales": len(month_sales),
            "total": month_total
        },
        "products": {
            "total": total_products,
            "low_stock": len(low_stock_products)
        },
        "customers": total_customers
    }

# ==================== REPORTS API ====================

@api_router.get("/reports/sales")
async def get_sales_report(
    period: str = Query(..., regex="^(daily|weekly|monthly)$"),
    date: Optional[str] = None
):
    """Get sales report for a specific period"""
    now = datetime.now(COLOMBIA_TZ) if not date else datetime.fromisoformat(date).replace(tzinfo=COLOMBIA_TZ)
    
    if period == "daily":
        start_date = now.replace(hour=0, minute=0, second=0, microsecond=0)
        end_date = start_date + timedelta(days=1)
    elif period == "weekly":
        start_date = now - timedelta(days=now.weekday())
        start_date = start_date.replace(hour=0, minute=0, second=0, microsecond=0)
        end_date = start_date + timedelta(days=7)
    else:  # monthly
        start_date = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        if now.month == 12:
            end_date = start_date.replace(year=now.year + 1, month=1)
        else:
            end_date = start_date.replace(month=now.month + 1)
    
    # Get sales in period
    sales = await db.sales.find({
        "created_at": {
            "$gte": datetime_to_str(start_date),
            "$lt": datetime_to_str(end_date)
        }
    }).to_list(10000)
    
    # Calculate statistics
    total_sales = len(sales)
    total_revenue = sum(sale['total'] for sale in sales)
    
    # Products sold
    products_sold = {}
    for sale in sales:
        for item in sale['items']:
            if item['product_id'] not in products_sold:
                products_sold[item['product_id']] = {
                    'name': item['product_name'],
                    'quantity': 0,
                    'revenue': 0
                }
            products_sold[item['product_id']]['quantity'] += item['quantity']
            products_sold[item['product_id']]['revenue'] += item['subtotal']
    
    # Sort by quantity
    top_products = sorted(
        products_sold.values(),
        key=lambda x: x['quantity'],
        reverse=True
    )[:10]
    
    return {
        "period": period,
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "total_sales": total_sales,
        "total_revenue": total_revenue,
        "average_sale": total_revenue / total_sales if total_sales > 0 else 0,
        "top_products": top_products
    }

@api_router.get("/reports/export")
async def export_report_to_excel(
    period: str = Query(..., regex="^(daily|weekly|monthly)$"),
    date: Optional[str] = None
):
    """Export sales report to Excel"""
    # Get report data
    report_data = await get_sales_report(period, date)
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {period.capitalize()}"
    
    # Styles
    header_fill = PatternFill(start_color="FFB6D7", end_color="FFB6D7", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Header
    ws['A1'] = "REPORTE DE VENTAS NOVAVENTA"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:D1')
    
    # Report info
    ws['A3'] = "Período:"
    ws['B3'] = period.capitalize()
    ws['A4'] = "Fecha Inicio:"
    ws['B4'] = report_data['start_date']
    ws['A5'] = "Fecha Fin:"
    ws['B5'] = report_data['end_date']
    
    # Summary
    ws['A7'] = "RESUMEN"
    ws['A7'].font = header_font
    ws['A7'].fill = header_fill
    ws.merge_cells('A7:B7')
    
    ws['A8'] = "Total Ventas:"
    ws['B8'] = report_data['total_sales']
    ws['A9'] = "Ingresos Totales:"
    ws['B9'] = f"${report_data['total_revenue']:,.0f} COP"
    ws['A10'] = "Venta Promedio:"
    ws['B10'] = f"${report_data['average_sale']:,.0f} COP"
    
    # Top Products
    ws['A12'] = "PRODUCTOS MÁS VENDIDOS"
    ws['A12'].font = header_font
    ws['A12'].fill = header_fill
    ws.merge_cells('A12:D12')
    
    headers = ['Producto', 'Cantidad', 'Ingresos', 'Promedio']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=13, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
    
    row = 14
    for product in report_data['top_products']:
        ws[f'A{row}'] = product['name']
        ws[f'B{row}'] = product['quantity']
        ws[f'C{row}'] = f"${product['revenue']:,.0f}"
        ws[f'D{row}'] = f"${product['revenue']/product['quantity']:,.0f}"
        row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    # Save to BytesIO
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    filename = f"reporte_{period}_{datetime.now(COLOMBIA_TZ).strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ==================== ROOT ENDPOINT ====================

@api_router.get("/")
async def root():
    return {"message": "NOVAVENTA API", "status": "active"}

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
